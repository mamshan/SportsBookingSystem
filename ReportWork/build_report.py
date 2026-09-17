from pathlib import Path
import json, hashlib, re, textwrap, zipfile
from datetime import datetime
from collections import defaultdict
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageChops

ROOT = Path(__file__).resolve().parent.parent
WORK = ROOT / 'ReportWork'
EVIDENCE = WORK / 'evidence'
SOURCE = ROOT.parent.parent / 'My Report.docx'
OUTPUT = ROOT / 'Report'
OUTPUT.mkdir(exist_ok=True)
db = json.loads((EVIDENCE / 'database.json').read_text())
doc = Document(SOURCE)
source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
section = doc.sections[0]
WORK.joinpath('artifact.md').write_text(
    f'# Report editing contract\nSource: {SOURCE}\nSHA256: {source_hash}\n'
    'Source: 25 pages, one A4 portrait section, 1-inch margins. Retain section geometry, Times New Roman body, black numbered headings and centered page footer.\n'
    'Edit scope: cover fields, TOC, all report content, tables and figures. Preserve original file. Correct inconsistent schema descriptions.\n'
    'Rebuild stale TOC from Heading 1-3 and refresh it in Word. Keep 16-point Heading 1 and 14-point Heading 2.\n'
    'Use real project diagrams and application screenshots. Tables describe actual SQL metadata. Missing personal details remain writable lines.\n', encoding='utf-8')
for element in list(doc._element.body):
    if element.tag != qn('w:sectPr'):
        doc._element.body.remove(element)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','List Bullet']:
    if name not in doc.styles:
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.space_after = Pt(6)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
for name,size in [('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles[name].paragraph_format.space_before = Pt(12)
doc.styles['Caption'].font.size=Pt(10)
doc.styles['Caption'].font.italic=True
for f in section.footer.paragraphs:
    f.clear()
p=section.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)

def para(text='',style=None):
    return doc.add_paragraph(text,style)
def heading(text,level=1):
    return doc.add_heading(text,level)
def page():
    doc.add_page_break()
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,value in zip(cells,row):c.text=str(value if value is not None else '')
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    pr.append(borders)
    for i,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for j,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[j])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'70');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');tcpr.append(sh)
            for p in c.paragraphs:
                if len(t.rows) <= 9:
                    p.paragraph_format.keep_with_next = i < len(t.rows)-1
                p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.line_spacing=1
                for r in p.runs:r.font.size=Pt(9);r.bold=(i==0)
    para()
    return t
def code(text):
    lines=text.strip().splitlines()
    for index,line in enumerate(lines):
        for wrapped in textwrap.wrap(line,width=94,replace_whitespace=False,drop_whitespace=False) or ['']:
            p=para()
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1
            p.paragraph_format.keep_with_next = index < len(lines)-1 and len(lines)<35
            r=p.add_run(wrapped);r.font.name='Consolas';r.font.size=Pt(8)
    para()
def figure(filename,caption,width=6.2):
    file=Path(filename)
    if not file.is_absolute():file=EVIDENCE/file
    im=Image.open(file)
    if file.parent == EVIDENCE:
        rgb=im.convert('RGB')
        box=ImageChops.difference(rgb,Image.new('RGB',rgb.size,'white')).point(lambda x:255 if x>20 else 0).getbbox()
        if box and box[3]+20 < im.height:
            im=im.crop((0,0,im.width,box[3]+20))
            crop=EVIDENCE/'cropped';crop.mkdir(exist_ok=True)
            file=crop/file.name;im.save(file)
    width=min(width,6.2)
    height=width*im.height/im.width
    if height>6.3:width=6.3*im.width/im.height
    p=para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(file),width=Inches(width))
    p=para(caption,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
def excerpt(file,start,end):
    lines=(ROOT/file).read_text(encoding='utf-8-sig').splitlines()
    code('\n'.join(lines[start-1:end]))
def sqltype(c):
    t=c['DataType'].upper()
    if t in ('VARCHAR','CHAR','NVARCHAR','NCHAR'):
        size=c['MaxLength']//2 if t.startswith('N') else c['MaxLength']
        t+=f'({size})'
    if t in ('DECIMAL','NUMERIC'):t+=f"({c['Precision']},{c['Scale']})"
    return t
columns=defaultdict(list)
for c in db['columns']:columns[c['TableName']].append(c)
order=['Member','FacilityType','Facility','Booking','Review','SportPreference','Inquiry']
ddl=[]
for name in order:
    fields=[]
    for c in columns[name]:
        fields.append('    '+c['ColumnName']+' '+sqltype(c)+(' IDENTITY(1,1)' if c['IsIdentity'] else '')+(' NULL' if c['Nullable'] else ' NOT NULL'))
    for keyname in dict.fromkeys(k['ConstraintName'] for k in db['keys'] if k['TableName']==name):
        ks=[k for k in db['keys'] if k['ConstraintName']==keyname]
        kind='PRIMARY KEY' if ks[0]['Kind']=='PRIMARY_KEY_CONSTRAINT' else 'UNIQUE'
        fields.append('    CONSTRAINT '+keyname+' '+kind+' ('+', '.join(k['ColumnName'] for k in ks)+')')
    for fk in db['foreignKeys']:
        if fk['TableName']==name:
            fields.append('    CONSTRAINT '+fk['ConstraintName']+' FOREIGN KEY ('+fk['ColumnName']+') REFERENCES '+fk['TargetTable']+'('+fk['TargetColumn']+')')
    for ck in db['checks']:
        if ck['TableName']==name:fields.append('    CONSTRAINT '+ck['ConstraintName']+' CHECK '+ck['Definition'])
    ddl.append('CREATE TABLE '+name+' (\n'+',\n'.join(fields)+'\n);\nGO')
OUTPUT.joinpath('DatabaseSchema.sql').write_text('-- Recreates the verified schema in an empty database.\n\n'+'\n\n'.join(ddl),encoding='utf-8')
OUTPUT.joinpath('SelectQueries.sql').write_text((WORK/'queries.sql').read_text(),encoding='utf-8')
sample="""SET XACT_ABORT ON;
BEGIN TRANSACTION;
INSERT INTO Member (Name,Email,Phone,Address,Password,RegDate)
VALUES ('Sample Member','coursework.sample@example.test','0770000000',
        'Negombo','Sample123!',CAST(GETDATE() AS DATE));
DECLARE @MemberID INT = SCOPE_IDENTITY();
INSERT INTO FacilityType (TypeName) VALUES ('Sample Tennis');
DECLARE @TypeID INT = SCOPE_IDENTITY();
INSERT INTO Facility (Name,TypeID,Location,Capacity,HourlyRate)
VALUES ('Sample Tennis Court',@TypeID,'Negombo',4,1500.00);
DECLARE @FacilityID INT = SCOPE_IDENTITY();
INSERT INTO SportPreference (MemberID,TypeID) VALUES (@MemberID,@TypeID);
INSERT INTO Booking (MemberID,FacilityID,BookingDate,StartTime,EndTime,Status)
VALUES (@MemberID,@FacilityID,'2026-08-01','2026-08-01T09:00:00',
        '2026-08-01T10:00:00','Confirmed');
INSERT INTO Review (MemberID,FacilityID,Rating,Comments,ReviewDate)
VALUES (@MemberID,@FacilityID,4,'The court was clean.','2026-08-01T11:00:00');
INSERT INTO Inquiry (GuestName,Email,Message,DateSent)
VALUES ('Sample Guest','guest@example.test','Are weekend sessions available?',
        CAST(GETDATE() AS DATE));
COMMIT;"""
OUTPUT.joinpath('SampleData.sql').write_text('-- Run once in a test database after DatabaseSchema.sql.\n'+sample,encoding='utf-8')

p=para('Community Sports Facilities Booking System','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before=Pt(80);p.paragraph_format.space_after=Pt(18)
p=para('Design, Implementation and Testing of a Web-Based Database Application','Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
para()
for label in ['Module code and name','Student name','Student ID','Module leader','Submission date']:
    p=para(label+': __________________________________');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('Assignment No: 001');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
para()
para('AI assistance declaration')
para('OpenAI Codex assisted with code changes, SQL checks, automated testing, screenshots, report drafting and formatting. The student is responsible for reviewing the work, understanding the implementation and following the institution’s rules on permitted AI use.')
para('Student signature: _________________________    Date: ______________')
page()
p=para('Table of Contents');p.runs[0].bold=True;p.runs[0].font.size=Pt(16)
p=para();start=OxmlElement('w:fldChar');start.set(qn('w:fldCharType'),'begin')
instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' TOC \\o "1-3" \\h \\z \\u '
sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate')
end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
for el in [start,instr,sep,end]:
    r=OxmlElement('w:r');r.append(el);p._p.append(r)
page()
heading('1. Introduction')
para('The Community Sports Facilities Booking System is a web application for a municipal sports council. It helps members find local venues and request a suitable time slot. The application checks existing bookings before saving a new request, so overlapping bookings can be rejected.')
para('Registered members can sign in, record preferred sports, search facilities by sport, location, date and time, request bookings and review facilities after a confirmed visit. Guests can browse general availability, read reviews, register and send inquiries. Guests cannot create bookings or submit reviews.')
para('The prototype uses Microsoft SQL Server, ASP.NET Core MVC with C#, Entity Framework Core and Razor views. The database contains Member, FacilityType, Facility, Booking, Review, SportPreference and Inquiry. SQL Developer Data Modeler was used for the supplied design. The report explains the database, implementation, SQL queries, testing and remaining limitations.')
heading('1.1 Scope and assumptions',2)
for t in ['A facility supports one booking for an overlapping time slot. Adjacent bookings are allowed when one ends exactly as another starts.',
          'New bookings have Pending status. Confirmation is a manual database task because an administrative approval page is outside the prototype.',
          'Reviews require a Confirmed booking whose end time is in the past. More than one review is currently allowed.',
          'All dates and times use the server’s local time. Searches containing only a date show facilities with no active booking that day.',
          'The prototype has no payment processing, email confirmation or password recovery. Passwords remain plain text and must be hashed before production use.']:
    para(t,'List Bullet')

page();heading('2. Specification of Database Relations')
heading('2.1 Relational Entity-Relationship Diagram',2)
figure(ROOT.parent.parent/'ERD.png','Figure 2.1. Relational model exported from SQL Developer Data Modeler.')
para('The relational model contains seven entities. The diagram is consistent with the implemented relationships, including Booking to Member and Booking to Facility. SportPreference has its own primary key and a unique MemberID–TypeID pair. Inquiry is independent because a guest does not need a Member record.')
entities={
'Member':'Stores each registered person’s name, email, contact details, password and registration date. Email is unique and is used for login.',
'FacilityType':'Stores sport categories. Facilities and member preferences refer to a TypeID instead of repeating category names.',
'Facility':'Stores venue name, sport category, location, capacity and hourly rate.',
'Booking':'Links one member to one facility and stores the requested date, start time, end time and booking status.',
'Review':'Links the author and facility and stores an integer rating from 1 to 5, comments and a review date.',
'SportPreference':'Resolves the many-to-many relationship between members and sport categories. A member can choose several sports, but cannot select the same category twice.',
'Inquiry':'Stores a guest’s name, email, message and date. It has no foreign key to Member.'}
for name,text in entities.items():para(name+'. '+text)
heading('2.1.1 Relationships and cardinality',3)
table(['Parent','Child','Cardinality'],[(fk['TargetTable'],fk['TableName'],'One to many') for fk in db['foreignKeys']],[1.7,2.3,2.2])
para('Every child foreign key is required. One parent can have no related child records or many of them; every child row refers to exactly one parent. SportPreference turns the many-to-many association into two one-to-many relationships.')
heading('2.1.2 Normalisation',3)
para('The design separates different subjects into their own tables. In first normal form, each column holds a single value, and preferred sports are represented by separate SportPreference rows. The non-key attributes describe their table’s key rather than only part of a composite candidate key, supporting second normal form. FacilityType holds the sport name once, while Facility and SportPreference hold its ID, avoiding a repeated type name and supporting third normal form.')
heading('2.2 Data Dictionary',2)
para('The following tables describe the inspected SQL Server schema after reconciliation. Existing VARCHAR(255) lengths and DECIMAL(25,2) rates are retained to avoid truncating saved data. Earlier narrower lengths in the draft did not describe the actual database. Blank optional values are permitted only where NULL is shown.')
descriptions={'MemberID':'Member identifier','FacilityID':'Facility identifier','TypeID':'Sport category identifier','BookingID':'Booking identifier','ReviewID':'Review identifier','SportPreferenceID':'Preference identifier','InquiryID':'Inquiry identifier','Name':'Name','Email':'Email address','Phone':'Contact number','Address':'Postal address','Password':'Prototype login password','RegDate':'Registration date','Location':'Venue location','Capacity':'Maximum participants','HourlyRate':'Rate in LKR per hour','TypeName':'Sport category','BookingDate':'Requested date','StartTime':'Start date and time','EndTime':'End date and time','Status':'Pending, Confirmed or Cancelled','Rating':'Whole number from 1 to 5','Comments':'Written feedback','ReviewDate':'Feedback date and time','GuestName':'Guest sender','Message':'Inquiry text','DateSent':'Inquiry date'}
for i,name in enumerate(order,1):
    heading(f'2.2.{i} {name}',3)
    rows=[]
    for c in columns[name]:
        cons=['NULL allowed' if c['Nullable'] else 'NOT NULL']
        if c['IsIdentity']:cons.append('IDENTITY')
        for k in db['keys']:
            if k['TableName']==name and k['ColumnName']==c['ColumnName']:
                cons.append('PK' if k['Kind']=='PRIMARY_KEY_CONSTRAINT' else ('Unique pair with TypeID' if name=='SportPreference' and c['ColumnName']=='MemberID' else 'Unique pair with MemberID' if name=='SportPreference' else 'UNIQUE'))
        for fk in db['foreignKeys']:
            if fk['TableName']==name and fk['ColumnName']==c['ColumnName']:cons.append('FK to '+fk['TargetTable'])
        if c['ColumnName']=='Rating':cons.append('CHECK 1–5')
        if name=='Booking' and c['ColumnName']=='Status':cons.append('CHECK allowed statuses')
        rows.append((c['ColumnName'],sqltype(c),'; '.join(cons),descriptions[c['ColumnName']]))
    table(['Attribute','SQL type','Key and constraints','Description'],rows,[1.05,1.15,2.15,1.85])
heading('2.3 SQL Developer Data Modeler',2)
para('The supplied design is saved as ER.dmd with its ER model folder. The original logical diagram was engineered into a relational model. The relational export in Figure 2.1 shows the final seven-table design and foreign keys. The saved DDL identifies SQL Developer Data Modeler 24.3.1.351.0831 and a SQL Server 2012 generation target. This is a DDL target setting, not a claim about the installed SQL Server version.')
code('\n'.join((ROOT.parent.parent/'ERDB File.ddl').read_text(encoding='utf-8-sig').splitlines()[:4]))
para('The generated script was checked against SQL Server metadata. The supplied DDL text omitted Booking_Member_FK even though the relational diagram showed it; the running database contains that foreign key. Review.Rating was still VARCHAR in the running database. It was converted to INT after checking every saved value, and database CHECK constraints were added for Rating and Status. Appendix A therefore records the reconciled database rather than repeating an outdated generated script.')
para('The original logical export is an earlier draft: it lacks SportPreferenceID and has relationship lines that do not fully match the physical schema. Figure 2.1 and the relationship table above define the implemented design. The saved project, original exports and reconciled DDL are retained as design evidence.')

page();heading('3. Database Implementation')
heading('3.1 Data Models',2)
para('SQL Server stores all seven tables. Primary keys use IDENTITY(1,1), foreign keys preserve references, Member.Email is unique and the SportPreference pair is unique. EF Core maps BookingID as database-generated and MemberID as a supplied foreign key. The previous mapping reversed these settings; the corrected mapping allows a new booking to receive its own identifier.')
code('entity.Property(e => e.BookingId).ValueGeneratedOnAdd();\nentity.Property(e => e.MemberId).ValueGeneratedNever();')
para('Tables can be created in dependency order: Member and FacilityType, then Facility, followed by Booking, Review and SportPreference. Inquiry has no dependencies. Appendix A contains the complete executable CREATE TABLE script extracted from the inspected schema.')
figure('constraints.png','Figure 3.1. Actual SQL Server CHECK constraint metadata, displayed as a result table.')
heading('3.2 Insert Statements',2)
para('The database already contains representative sample records. The captured counts below exclude the temporary records created for testing and screenshots. Existing data includes duplicate facility names with different IDs, so relationships and queries use IDs rather than assuming names are unique.')
table(['Table','Saved rows'],[(r['TableName'],r['Rows']) for r in db['counts']],[4.4,1.8])
para('Appendix B provides a standalone example that inserts into every table in dependency order. SCOPE_IDENTITY() captures each new key instead of assuming IDs begin at 1. The example is intended to be run once in a test database; it was not added to the existing database. Successful application inserts are demonstrated by the captured Member, SportPreference, Booking, Review and Inquiry rows in Section 4.')
code("INSERT INTO Facility (Name, TypeID, Location, Capacity, HourlyRate)\nVALUES ('Sample Tennis Court', @TypeID, 'Negombo', 4, 1500.00);\nDECLARE @FacilityID INT = SCOPE_IDENTITY();")
heading('3.3 Select Queries',2)
para('Eight queries were executed against SportsBooking. Each result image displays the actual returned rows in a readable table; these are captured query outputs, not screenshots of SQL Server Management Studio. No passwords are selected. The availability example uses 17 October 2026, 10:00–11:00.')
purposes=[
'This three-table join replaces raw foreign keys with readable member and facility names for a booking list.',
'The NOT EXISTS condition excludes overlapping Pending or Confirmed bookings. Cancelled bookings do not block a slot. Missing times are treated as blocking the requested day.',
'GROUP BY produces one row per reviewed facility. Casting before AVG prevents integer division from hiding fractional averages.',
'TOP and ORDER BY list five venues with the lowest known hourly rates. FacilityID breaks equal-price ties.',
'The subquery finds members with at least one booking, without duplicating a member who has several bookings.',
'LEFT JOIN retains facilities even when there is no matching booking. Testing for a null booking ID identifies unused facilities.',
'The link table joins members to their selected sport categories. This query demonstrates the many-to-many association.',
'The latest five inquiries are returned first, with InquiryID resolving ties in DateSent. This is an internal SQL report; the application does not expose inquiry records publicly.'
]
for i,(q,purpose) in enumerate(zip(db['queries'],purposes),1):
    page();heading(f'3.3.{i} {q["title"]}',3);para(purpose)
    formatted=q['sql']
    for token in [' FROM ',' WHERE ',' GROUP BY ',' ORDER BY ',' LEFT JOIN ',' JOIN ']:
        formatted=formatted.replace(token,'\n'+token.strip()+' ')
    code(formatted)
    figure(f'query-{i}.png',f'Figure 3.{i+1}. Query {i} returned {len(q["rows"])} rows.')

page();heading('4. Implementation of the Web-Based Application')
heading('4.1 System Architecture',2)
para('The project uses a simple MVC structure. Models describe the database entities and form inputs, controllers process requests, and Razor views render HTML. SportsContext is injected into controllers and uses the SQL Server provider to query and save data. EF Core reverse engineering produces entity classes and a DbContext from an existing database (Microsoft, n.d.-b).')
code('Browser request\n    -> Controller action\n    -> Form view model and validation\n    -> SportsContext and EF Core\n    -> SQL Server\n    -> Razor view returned to the browser')
table(['Area','Files and purpose'],[
('Controllers','Home, Account, Facility, Booking, Review and Inquiry handle public/member routes.'),
('Models','Seven entity classes; RegisterViewModel, FacilitySearchViewModel, BookingViewModel and ReviewViewModel handle form data.'),
('Data','SportsContext contains table mappings and relationships.'),
('Views','Razor pages share _Layout.cshtml and Bootstrap styling.'),
('Program.cs','Registers MVC, SQL Server and sessions; serves static files and maps routes.'),
('tests','SmokeTest.ps1 executes HTTP requests and checks saved database records.')],[1.2,5])
para('MVC separates request handling from page presentation (Microsoft, n.d.-a). Session stores MemberID and MemberName after a valid login. Protected actions check MemberID before accessing data. AddDistributedMemoryCache, AddSession and UseSession configure this prototype’s in-memory sessions (Microsoft, n.d.-c). Restarting the app clears them.')
excerpt('Program.cs',7,17)
para('View models contain only editable fields. MemberId and booking Status are set by the server. Required fields, maximum lengths and numeric ranges are checked through data annotations and ModelState.IsValid (Microsoft, n.d.-d). POST forms include anti-forgery tokens.')
heading('4.2 Home Page',2)
para('The home page introduces community sports, links to facility search and reviews, and presents separate member-login and new-registration areas. Navigation changes after login to include My Bookings and logout. Contact Us remains available to both members and guests.')
figure('home-guest.png','Figure 4.1. Home page with login and registration options.')
code('@if (Context.Session.GetInt32("MemberID") == null)\n{\n    <a asp-controller="Account" asp-action="Login">Member login</a>\n    <a asp-controller="Account" asp-action="Register">Register</a>\n}')
heading('4.3 Member Functionality',2)
heading('4.3.1 Sign In',3)
para('AccountController checks that email and password have been supplied, looks up the member and displays an error when the details do not match. A valid login stores the member identity in session and redirects to Facility/Search. Logout is a POST action that clears the session.')
excerpt('Controllers/AccountController.cs',28,50)
figure('login-error.png','Figure 4.2. Invalid credentials produce a validation message.')
para('The existing prototype compares plain-text passwords. This is a recorded limitation, not a production authentication design. Password hashing, account lockout and recovery are future improvements.')
heading('4.3.2 Registration',3)
para('The registration form captures name, email, phone, address, password and at least one preferred sport. It rejects duplicate email addresses and sport IDs that do not exist. Repeated sport IDs are removed with Distinct(). The Member and its SportPreferences are saved in one SaveChangesAsync call, then the member is signed in.')
figure('registration.png','Figure 4.3. Completed registration form with a preferred sport.')
code('foreach (var typeId in model.SelectedSports)\n{\n    member.SportPreferences.Add(new SportPreference { TypeId = typeId });\n}\n_context.Members.Add(member);\nawait _context.SaveChangesAsync();')
figure('data-member.png','Figure 4.4. Saved test member row, with the password excluded.')
figure('data-preference.png','Figure 4.5. Preferred sport saved for the same member.')
heading('4.3.3 Facility Search',3)
para('Members can combine sport, location, date and time filters. Empty filters are ignored. Both times are required when a time filter is used, and an end time must follow the start time. A query excludes any facility with an overlapping non-cancelled booking. The page shows no-results feedback when no venue matches.')
excerpt('Controllers/FacilityController.cs',25,38)
figure('search-member.png','Figure 4.6. Member search with availability, price and booking link.')
figure('search-unavailable.png','Figure 4.7. The same test facility is excluded when its slot is booked.')
heading('4.3.4 Facility Booking',3)
para('The form accepts a facility, date, start time and end time. MemberID comes from the session and cannot be chosen through the form. The controller rejects past times, invalid intervals, nonexistent facilities and clashes. New requests are saved as Pending. The My Bookings query filters by the logged-in member, so it does not list another member’s bookings.')
figure('booking-form.png','Figure 4.8. Booking form before submission.')
code('await using var transaction = await _context.Database\n    .BeginTransactionAsync(IsolationLevel.Serializable);\nvar clash = await _context.Bookings.AnyAsync(b =>\n    b.FacilityId == model.FacilityId && b.BookingDate == model.BookingDate\n    && b.Status != "Cancelled"\n    && (b.StartTime == null || b.EndTime == null\n        || (b.StartTime < end && b.EndTime > start)));')
para('The serializable transaction keeps the availability check and insert together. The overlap rule uses start < requested end and end > requested start. Therefore 10:00–11:00 and 11:00–12:00 do not overlap. Load testing and deadlock retry handling are not included in the current prototype.')
figure('booking-saved.png','Figure 4.9. Saved booking displayed in My Bookings.')
figure('data-booking.png','Figure 4.10. Saved booking rows used for screenshot evidence. The 2020 Confirmed row was deliberately seeded to test reviews.')
heading('4.3.5 Review Submission',3)
para('A member can review only a facility with a Confirmed booking that has already ended. The form provides eligible facilities, an integer rating from 1 to 5 and optional comments up to 255 characters. The server supplies MemberID and ReviewDate. The database also enforces the rating range through CK_Review_Rating.')
code('if (!await _context.Bookings.AnyAsync(b =>\n    b.MemberId == memberId && b.FacilityId == model.FacilityId\n    && b.Status == "Confirmed" && b.EndTime < DateTime.Now))\n{\n    ModelState.AddModelError("FacilityId",\n        "You can review a facility after a confirmed booking has ended.");\n}')
figure('review-form.png','Figure 4.11. Review form for an eligible facility.')
figure('data-review.png','Figure 4.12. Saved review with its integer rating and date.')
heading('4.4 Guest Functionality',2)
heading('4.4.1 Restricted Search',3)
para('Guests can use the search filters to see facility names, sport types, locations and broad availability. They do not see member-only booking links, capacity or hourly-rate columns. This display rule is backed by server checks on booking and review creation. Direct requests to those actions redirect guests to login.')
figure('search-guest.png','Figure 4.13. Guest search without member-only booking controls.')
code('var memberId = HttpContext.Session.GetInt32("MemberID");\nif (memberId == null)\n    return RedirectToAction("Login", "Account");')
para('Unused scaffolded Member, FacilityType and SportPreference controllers are marked NonController. Facility, Booking, Review and Inquiry controllers expose only the intended routes. Old management URLs return 404; hiding links alone would not protect the underlying actions.')
heading('4.4.2 Review Search',3)
para('Guests can read reviews without an account. Optional filters select a facility and/or rating. The query includes Facility and Member so the page shows meaningful names instead of foreign-key numbers. The public page has no edit or delete links.')
figure('review-guest.png','Figure 4.14. A guest reads the submitted four-star review.')
code('var reviews = _context.Reviews\n    .Include(r => r.Facility).Include(r => r.Member).AsQueryable();\nif (facilityId.HasValue)\n    reviews = reviews.Where(r => r.FacilityId == facilityId);\nif (rating.HasValue)\n    reviews = reviews.Where(r => r.Rating == rating.Value);')
heading('4.4.3 Membership Registration',3)
para('The home page and guest navigation both link to Account/Register. Guests use the same registration form described in Section 4.3.2, which includes the required name and email as well as contact details and sport preferences. After a successful submission the guest becomes a member and reaches the member search page.')
heading('4.4.4 Send Inquiry',3)
para('Inquiry/Create is public. Name, a valid email and message are required. The application sets DateSent, saves the row and displays a confirmation. Inquiry does not need a MemberID. The application no longer exposes a public list of other people’s inquiries.')
figure('inquiry-form.png','Figure 4.15. Guest inquiry form.')
code('if (!ModelState.IsValid) return View(inquiry);\ninquiry.DateSent = DateOnly.FromDateTime(DateTime.Today);\n_context.Inquiries.Add(inquiry);\nawait _context.SaveChangesAsync();\nTempData["Success"] = "Your inquiry has been sent to the Sports Council.";')
figure('inquiry-saved.png','Figure 4.16. Confirmation after saving an inquiry.')
figure('data-inquiry.png','Figure 4.17. Corresponding Inquiry row in SQL Server.')

page();heading('5. Testing')
para('Testing used the running application and local SQL Server database on 17 September 2026. The PowerShell script tests HTTP routes, validation and persistence. A separate headless Microsoft Edge run checks rendered pages and captures screenshots. Both use uniquely identified temporary records and remove them afterwards. The screenshots show deliberate test data, including a past Confirmed booking needed for review eligibility.')
tests=[
('Home page','GET /','Sports information visible','Home heading returned'),
('Guest facility search','Type filter; no login','Browse without rate or booking details','Facility visible; rate column absent'),
('Guest booking form','GET /Booking/Create','Redirect to login','Login redirect'),
('Guest booking list','GET /Booking/MyBookings','Redirect to login','Login redirect'),
('Guest review form','GET /Review/Create','Redirect to login','Login redirect'),
('Member management','GET /Member','Route disabled','404'),
('Facility creation','GET /Facility/Create','Route disabled','404'),
('Facility editing','GET /Facility/Edit/1','Route disabled','404'),
('Review deletion','GET /Review/Delete/1','Route disabled','404'),
('Inquiry list','GET /Inquiry','Route disabled','404'),
('Preference management','GET /SportPreference','Route disabled','404'),
('Sport management','GET /FacilityType','Route disabled','404'),
('Invalid login','Unknown email, wrong password','Display error','Invalid credentials message'),
('Invalid sport','SelectedSports=999999','Reject unknown sport','Validation message'),
('Registration','Valid new details and sport','Save and sign in','Redirect to search'),
('Preference insert','New member and sport IDs','One matching row','SQL count = 1'),
('Duplicate email','Register same email again','Reject duplicate','Already registered message'),
('Valid login','New test email, correct password','Sign in','Redirect to search'),
('Member search','Logged-in type search','Show rate and Book link','Both shown'),
('Booking save','Future date, 10:00–11:00','Save booking','Confirmation shown'),
('Forged booking fields','MemberId=999999, Status=Confirmed','Use session ID and Pending','SQL verified server values'),
('Overlap','Existing 10:00–11:00; request 10:30–11:30','Reject clash','Already booked message'),
('Availability search','Search booked 10:30–11:30','Exclude booked facility','No results'),
('Adjacent slot','Request 11:00–12:00','Allow booking','Confirmation shown'),
('Invalid interval','Start 12:00, end 11:00','Reject interval','End-time validation'),
('Past booking','1 January 2020','Reject past time','Future-time validation'),
('Review before visit','Only future Pending booking','Reject review','Eligibility validation'),
('Invalid rating','Eligible visit, rating 6','Reject out-of-range value','Range validation'),
('Valid review','Eligible visit, rating 4','Save review','Confirmation shown'),
('Review filter','Facility and rating=4','Matching review visible','Expected content returned'),
('Empty inquiry','Empty Message','Reject required field','Required-field validation'),
('Inquiry submit','Valid name, email and message','Save and confirm','Confirmation shown'),
('Inquiry persistence','Query saved inquiry','One row with today’s date','SQL count = 1'),
('Logout','POST /Account/Logout','Clear session','Protected page redirects')
]
passed_lines=[s for s in (EVIDENCE/'smoke-results.txt').read_text(encoding='utf-8-sig').splitlines() if s.startswith('PASS:')]
assert len(passed_lines)==len(tests)
table(['ID','Test and data','Expected result','Actual result','Status'],
      [(f'T{i:02}',a+'\n'+b,c,d,'Pass') for i,(a,b,c,d) in enumerate(tests,1)],[.42,2.0,1.45,1.85,.48])
heading('5.1 Test Evidence',2)
para('Figures 4.2–4.17 support the login, registration, booking, review and inquiry tests. The image below records the overlap validation from T22. The saved test script contains the repeatable requests and SQL assertions. The browser run separately confirms guest review access using a fresh, unauthenticated browser context.')
figure('booking-clash.png','Figure 5.1. T22: an overlapping booking produces an error.')
table(['Browser check','Result'],[(r['name'],'Pass' if r['pass'] else 'Fail') for r in json.loads((EVIDENCE/'browser-tests.json').read_text())],[5.3,.9])
heading('5.2 Summary of Testing',2)
para('All 34 automated HTTP/database checks passed. All nine additional browser checks passed. The project built with zero warnings and zero errors. These results cover the tested prototype flows; they do not establish production security, performance under load or complete browser compatibility.')
table(['Fault found','Correction','Retest'],[
('Search and My Bookings links targeted missing actions.','Implemented Facility/Search and Booking/MyBookings.','T15, T18–T20 passed.'),
('Scaffolded booking allowed posted identity/status fields.','Used a form view model; set identity and Pending status on the server.','T21 passed.'),
('Booking ID mapping did not match SQL identity columns.','BookingID generated on add; MemberID supplied normally.','T20 inserted a booking successfully.'),
('Review rating and report disagreed; CHECK constraints were absent.','Converted valid ratings to INT, updated EF and added both checks.','T28–T29 passed; database metadata verified.'),
('CSS did not load when running without a launch profile.','Added UseStaticFiles before routing.','Bootstrap resource returned HTTP 200; screenshots recaptured.'),
('Unused scaffolded routes exposed management operations.','Disabled unused controllers and removed public CRUD actions.','T06–T12 returned 404.')],[2.0,2.35,1.85])
para('Remaining limits include plain-text passwords, no automated booking approval, repeated reviews, no pagination and no concurrent-load test. The booking transaction uses serializable isolation, but a production version should handle deadlocks and retries. Form checks reject invalid user inputs; further database checks for times, rates and capacity would strengthen direct SQL writes.')

page();heading('6. Further Discussion and Reflection')
para('The main design decision in this project was to keep SQL Server as the source of the data model. This made it possible to compare the database diagram, SQL script, entity classes and report directly. It also exposed a weakness in relying on generated files without checking them: the Booking key mapping and Review rating type did not match the intended design.')
para('For this prototype, I kept the controllers and Razor pages simple. Each form has a small view model and the controller handles its validation and database operation. This makes it easier to follow the booking flow from the submitted date and time to the final Booking row. A larger application would need more separation of business rules and reusable authorization checks.')
para('The booking overlap rule was an important part of the implementation. Comparing only the start times would miss partially overlapping reservations. The implemented rule compares both ends of each interval. Tests confirmed that an overlapping request is rejected while an adjacent booking is accepted. The server also supplies MemberID and Status, which prevents a member from changing them by editing form data.')
para('The review flow links feedback to a completed Confirmed booking. This makes the meaning of a review clearer, but it also exposes a limitation in the workflow: the prototype has no approval screen. New requests stay Pending until the database is updated manually. An administrative approval page is therefore a useful next feature.')
para('Testing revealed that a successful build does not prove the whole page works. The CSS problem appeared only when the application was opened for screenshots without a launch profile. The application also needed checks that direct URLs were blocked for guests, because removing a navigation link alone does not protect an action.')
para('The project would need several improvements before public use. Passwords should be hashed, session-based identity checks should be replaced or strengthened with standard authentication, booking approvals should be managed through a protected interface, and confirmations could be sent by email. Search paging and broader concurrency tests would help with larger data sets. The practical lesson from the verified changes is to keep the schema, code, screenshots and written claims consistent, and to support each main feature with a repeatable test.')

heading('7. Conclusion')
para('The prototype provides member registration, login, facility search, booking requests, reviews and guest inquiries using ASP.NET Core MVC and SQL Server. The seven-table schema supports the required relationships, and the rating and status rules are enforced in the database. The tested application passed 34 HTTP/database checks and nine browser checks. It remains a coursework prototype with the limitations discussed above.')

heading('8. References')
refs=[
('Microsoft (n.d.-a). Overview of ASP.NET Core MVC.','https://learn.microsoft.com/en-us/aspnet/core/mvc/overview?view=aspnetcore-10.0'),
('Microsoft (n.d.-b). Reverse Engineering. Entity Framework Core.','https://learn.microsoft.com/en-us/ef/core/managing-schemas/scaffolding/'),
('Microsoft (n.d.-c). Session in ASP.NET Core.','https://learn.microsoft.com/en-us/aspnet/core/fundamentals/app-state?view=aspnetcore-10.0'),
('Microsoft (n.d.-d). Model validation in ASP.NET Core MVC.','https://learn.microsoft.com/en-us/aspnet/core/mvc/models/validation?view=aspnetcore-10.0'),
('Oracle (n.d.). Data Modeler Concepts and Usage. SQL Developer Data Modeler 24.3 User’s Guide.','https://docs.oracle.com/en/database/oracle/sql-developer-data-modeler/24.3/dmdug/data-modeler-concepts-usage.html')]
for title,url in refs:
    para(title+' Available at: '+url+' (Accessed: 17 September 2026).')
para('Coursework brief (undated). Community Sports Facilities Booking System. Supplied assignment document.')
para('Project evidence: SportsBookingSystem source files, ER.dmd, ERD.png, Logical.png, ERDB File.ddl, SQL metadata and query results, SmokeTest.ps1 output and browser screenshots captured on 17 September 2026.')

page();heading('9. Appendices')
heading('Appendix A Full database schema',2)
para('This script recreates the verified schema in an empty database. It is a schema export, not an instruction to drop or recreate the existing SportsBooking database.')
for name,script in zip(order,ddl):
    heading(name,3);code(script)
heading('Appendix B Sample insert script',2)
para('This example inserts one related set of sample records in a test database. The existing database’s wider sample data and actual row counts are documented in Section 3.2.')
code(sample)
heading('Appendix C Running and verifying the project',2)
para('Open SportsBookingSystem.csproj in Visual Studio or use the commands below. The configured database is SportsBooking on 127.0.0.1 with Windows authentication. SQL Server must be running, and the Windows account must have access.')
code('dotnet build\ndotnet run --no-launch-profile --urls http://127.0.0.1:5187\n# Run in a second PowerShell window.\n.\\tests\\SmokeTest.ps1')
para('The test script creates and removes only its uniquely named temporary member, facility, category, bookings, reviews and inquiry records. Its SQL cleanup runs in finally. It should be run against a local coursework/test database.')
heading('Appendix D Full controller code',2)
para('The complete controller source is supplied with the project. The main report includes the relevant excerpts to keep the explanations readable. Files: Controllers/AccountController.cs, FacilityController.cs, BookingController.cs, ReviewController.cs and InquiryController.cs. Program.cs contains service and middleware setup.')
settings=doc.settings.element
upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');settings.append(upd)
final=OUTPUT/'My Report Completed.docx'
doc.save(final)
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
print(final)
print('Tables:',len(doc.tables),'Figures:',len(doc.inline_shapes))
